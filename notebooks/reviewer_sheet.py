"""
Builds the physician blind-review workbook for Level-1 label validation
(architecture doc section 6): "ο γιατρός της ομάδας κάνει τυφλή αξιολόγηση
και στους 117 ασθενείς ... με φόρμα που ακολουθεί την ιεραρχία της
ενότητας 3."

Design choice: the physician reads the RAW note text (not our NLP-extracted
snippets) and fills in an independent determination. The pipeline's own
proposed labels are written to a SEPARATE sheet that should not be opened
until after the blind pass is complete — otherwise the "independent"
review is anchored on what the NLP system already highlighted, which would
inflate the apparent precision/recall/kappa computed in section 6 rather
than actually testing it.
"""
import pandas as pd


REVIEW_COLUMNS = [
    "Profile Key", "N_notes", "Notes_types_and_years",
    "SVD_present (Y/N)", "HVD_stage (0/1/2/3)", "HVD_phenotype (S/R/RS)",
    "BVF_stage (0/2/3)", "Confidence (definite/probable/possible)",
    "Exclusion_or_competing_event (endocarditis/thrombosis/other, or blank)",
    "Reviewer_notes",
]


def build_reviewer_workbook(notes_df: pd.DataFrame, labels_df: pd.DataFrame, out_path: str):
    patients = sorted(notes_df["Profile Key"].unique())

    # --- Sheet 1: raw notes, one row per note, for the physician to read ---
    notes_sheet = notes_df[["Profile Key", "Type", "Service Date", "Notes"]].copy()
    notes_sheet = notes_sheet.sort_values(["Profile Key", "Service Date"])
    notes_sheet.columns = ["Profile Key", "Note Type", "Service Date (year)", "Note Text"]

    # --- Sheet 2: blind determination form, one row per patient ----------
    form_rows = []
    for pid in patients:
        sub = notes_df[notes_df["Profile Key"] == pid].sort_values("Service Date")
        types_years = "; ".join(f"{t} {y}" for t, y in zip(sub["Type"], sub["Service Date"]))
        form_rows.append({
            "Profile Key": pid, "N_notes": len(sub), "Notes_types_and_years": types_years,
            "SVD_present (Y/N)": "", "HVD_stage (0/1/2/3)": "", "HVD_phenotype (S/R/RS)": "",
            "BVF_stage (0/2/3)": "", "Confidence (definite/probable/possible)": "",
            "Exclusion_or_competing_event (endocarditis/thrombosis/other, or blank)": "",
            "Reviewer_notes": "",
        })
    form_df = pd.DataFrame(form_rows, columns=REVIEW_COLUMNS)

    # --- Sheet 3: pipeline's own proposed labels (do not open before blind
    # review is complete — see module docstring) --------------------------
    pred_cols = ["profile_key", "index_implant_year", "index_implant_source",
                 "index_valve_model", "index_valve_size_mm", "hvd_stage", "bvf_stage",
                 "phenotype", "confidence_tier", "rationale", "redo_evidence",
                 "svd_explicit_evidence", "exclusion_reason", "ar_evidence",
                 "morphology_evidence"]
    pred_df = labels_df[pred_cols].copy()

    instructions = pd.DataFrame({"Instructions": [
        "Level-1 label validation (blind physician review).",
        "",
        "1. Do NOT open the 'NLP_predictions (open after)' sheet before completing your review.",
        "2. For each Profile Key in 'Blind review form', read all of that patient's notes in 'Raw notes'.",
        "3. Fill in SVD_present, HVD_stage, HVD_phenotype (S/R/RS per VARC-3 Table 13), BVF_stage, "
        "Confidence, and any exclusion/competing event, following the ground-truth hierarchy in "
        "architecture section 3 (Definite / Probable / Possible / Exclusion / Censored).",
        "4. Once complete, save and compare against 'NLP_predictions (open after)' to compute "
        "Cohen's kappa (SVD yes/no) and weighted kappa (stage), and precision/recall/F1 for the "
        "reintervention extraction, per architecture section 6 Level 1.",
        "5. A second reviewer should independently score a random 30-patient subset for inter-rater "
        "reliability (target kappa >= 0.7), also blind to the NLP predictions and to the first "
        "reviewer's answers.",
    ]})

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        instructions.to_excel(writer, sheet_name="Instructions", index=False)
        form_df.to_excel(writer, sheet_name="Blind review form", index=False)
        notes_sheet.to_excel(writer, sheet_name="Raw notes", index=False)
        pred_df.to_excel(writer, sheet_name="NLP_predictions (open after)", index=False)

    _style_workbook(out_path)
    return out_path


def _style_workbook(path: str):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    warn_fill = PatternFill("solid", fgColor="FFF2CC")

    for name in wb.sheetnames:
        ws = wb[name]
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            header = ws.cell(row=1, column=col_idx).value or ""
            width = 60 if ("Note" in str(header) or "rationale" in str(header) or
                            "evidence" in str(header) or "Instructions" in str(header)) else 22
            ws.column_dimensions[letter].width = width
        if name == "NLP_predictions (open after)":
            for row in ws.iter_rows(min_row=1, max_row=1):
                for cell in row:
                    cell.fill = warn_fill

    wb.save(path)


if __name__ == "__main__":
    import pandas as pd
    notes_df = pd.read_excel("notes_deidentified.xlsx")
    notes_df["Notes"] = notes_df["Notes"].astype(str)
    labels_df = pd.read_csv("labels.csv")
    path = build_reviewer_workbook(notes_df, labels_df, "physician_review_worksheet.xlsx")
    print("wrote", path)
